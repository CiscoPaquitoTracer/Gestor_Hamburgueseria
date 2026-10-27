"""
Convierte Comandas.xlsx (ventas de antes de la app) en SQL para la tabla historial_ventas,
que usa el pronóstico de ventas e insumos.

El Excel NO tiene fechas: cada hoja (Junio, Julio) tiene bloques de 7 columnas, uno por día.
Los bloques se leen en orden (renglón por renglón, de izquierda a derecha) y se asignan a los
días que abren, seguidos: con --dias vie,sab,dom el 1.º bloque es el primer viernes del mes,
el 2.º el sábado, el 3.º el domingo, el 4.º el viernes siguiente… Si se sabe la fecha del
primer bloque, se da con --inicio.

Uso (requiere openpyxl; el Excel y el SQL generado NO se suben al repo, que es público):
    python importar_comandas.py Comandas.xlsx --anio 2026 --dias vie,sab,dom > comandas.sql
    python importar_comandas.py Comandas.xlsx --anio 2026 --dias vie,sab,dom --inicio Junio=2026-06-05

Después: aplicar comandas.sql en Supabase. Se puede volver a correr: primero borra lo
importado antes desde este mismo archivo.
"""
import argparse
import datetime as dt
import sys
import unicodedata

import openpyxl

FUENTE = "Comandas.xlsx"
MESES = {"Junio": 6, "Julio": 7}
DIAS = {"lun": 0, "mar": 1, "mie": 2, "jue": 3, "vie": 4, "sab": 5, "dom": 6}

# Categoría del Excel → (categoría en Supabase, productos en Supabase).
MENU = {
    "HAMBURGUESAS": ("Burgers", ["Junior", "Glotona", "La Guacamole", "Sweet Bacon", "La Tóxica", "Hula Hula", "Bonchosa"]),
    "DOGGOS": ("Doggos Jumbo", ["El Perro Callejero", "Atascado", "Junior", "Pica Pica", "Mexicanísimo", "Tóxico Guac"]),
    "PAPAS": ("Papas", ["Normalitas", "Monchosas", "Junior", "Criminales", "BBQ Bacon"]),
    "REFRESCOS": ("Refrescos", ["Boing 500 ml", "Coca Cola 600 ml", "Sprite 600 ml", "Jarritos 600 ml", "Fresca 600 ml"]),
    "COMBOS": ("Combos", ["Solitario", "Desmadre de Pareja", "Caja de Pandora", "Kit Para El Bajón"]),
}
# Se venden por tamaño: el renglón del Excel es la variante.
CON_VARIANTE = {"ALITAS": ("Alitas", "Alitas"), "BONELESS": ("Boneless", "Boneless")}
NO_SON_PRODUCTOS = {"TOTAL", "TOTAL PEDIDOS", "EFECTIVO", "TRANSFERENCIA", "TARJETA", "TOTAL VENTAS", "MERCADO"}


def normal(texto: str) -> str:
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return " ".join(sin_acentos.lower().split())


def producto_db(categoria: str, nombre: str):
    """(categoría, producto, variante) en Supabase para un renglón del Excel."""
    if categoria in CON_VARIANTE:
        cat, prod = CON_VARIANTE[categoria]
        return cat, prod, nombre.strip()                  # "5 piezas"
    cat, productos = MENU[categoria]
    for p in productos:
        # "Coca" → "Coca Cola 600 ml", "La guacamole" → "La Guacamole"
        if normal(p) == normal(nombre) or normal(p).startswith(normal(nombre) + " "):
            return cat, p, None
    raise SystemExit(f"No reconozco '{nombre}' en {categoria}: agrégalo a MENU")


def leer_bloque(ws, r0: int, c0: int):
    """{(categoría, producto, variante): cantidad} de un día."""
    vendidos, categoria = {}, None
    for r in range(r0 + 1, ws.max_row + 1):
        nombre, cantidad = ws.cell(r, c0).value, ws.cell(r, c0 + 1).value
        if nombre == "Menú":
            break
        if not isinstance(nombre, str) or nombre.strip() in NO_SON_PRODUCTOS:
            continue
        if nombre.strip() in MENU or nombre.strip() in CON_VARIANTE:
            categoria = nombre.strip()
        elif isinstance(cantidad, (int, float)) and cantidad > 0:
            clave = producto_db(categoria, nombre)
            vendidos[clave] = vendidos.get(clave, 0) + cantidad
    return vendidos


def fechas_de_hoja(ws, anio: int, mes: int, dias: list, inicio):
    """[(fecha, r0, c0)]: los bloques, en orden de lectura, van a los días que abren seguidos."""
    bloques = sorted((c.row, c.column) for fila in ws.iter_rows() for c in fila if c.value == "Menú")
    abiertos = {DIAS[d] for d in dias}
    fecha = inicio
    if fecha is None:   # primer día de apertura del mes (ej. primer viernes)
        fecha = dt.date(anio, mes, 1)
        while fecha.weekday() != DIAS[dias[0]]:
            fecha += dt.timedelta(days=1)
    if fecha.weekday() not in abiertos:
        raise SystemExit(f"{ws.title}: {fecha} no es un día que abran ({','.join(dias)})")
    resultado = []
    for r0, c0 in bloques:
        resultado.append((fecha, r0, c0))
        fecha += dt.timedelta(days=1)
        while fecha.weekday() not in abiertos:
            fecha += dt.timedelta(days=1)
    return resultado


def sql_texto(v):
    return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel")
    ap.add_argument("--anio", type=int, required=True)
    ap.add_argument("--dias", required=True, help="días que abrían, en orden de columnas: ej. mie,jue,vie,sab")
    ap.add_argument("--inicio", action="append", default=[], help="Hoja=AAAA-MM-DD del primer bloque (opcional)")
    args = ap.parse_args()

    dias = args.dias.split(",")
    inicios = {h: dt.date.fromisoformat(f) for h, f in (x.split("=") for x in args.inicio)}
    wb = openpyxl.load_workbook(args.excel, data_only=True)

    renglones, resumen = [], []
    for hoja, mes in MESES.items():
        for fecha, r0, c0 in fechas_de_hoja(wb[hoja], args.anio, mes, dias, inicios.get(hoja)):
            vendidos = leer_bloque(wb[hoja], r0, c0)
            resumen.append(f"-- {hoja} {openpyxl.utils.get_column_letter(c0)}{r0} → {fecha:%a %Y-%m-%d}: "
                           f"{sum(vendidos.values()):g} productos")
            renglones += [(fecha, *clave, cantidad) for clave, cantidad in vendidos.items()]

    fechas = [r[0] for r in renglones]
    if len(set((r[0], r[1], r[2], r[3]) for r in renglones)) != len(renglones):
        raise SystemExit("Dos bloques cayeron en la misma fecha: revisa --dias / --inicio")

    print("\n".join(resumen), file=sys.stderr)
    valores = ",\n".join(
        f"    ('{f}', {sql_texto(c)}, {sql_texto(p)}, {sql_texto(v)}, {n})" for f, c, p, v, n in renglones)
    print(f"""-- Generado por supabase/importar_comandas.py ({len(renglones)} renglones, {min(fechas)} a {max(fechas)}).
BEGIN;
DELETE FROM historial_ventas WHERE fuente = '{FUENTE}';

INSERT INTO historial_ventas (dia, producto_id, variante_id, cantidad, componente, fuente)
SELECT x.dia::date, p.id, v.id, x.cantidad, false, '{FUENTE}'
  FROM (VALUES
{valores}
  ) AS x (dia, categoria, producto, variante, cantidad)
  JOIN categorias_menu c ON c.nombre = x.categoria
  JOIN productos p       ON p.categoria_id = c.id AND p.nombre = x.producto
  LEFT JOIN variantes v  ON v.producto_id = p.id AND v.nombre = x.variante;

DO $$ BEGIN
    IF (SELECT count(*) FROM historial_ventas WHERE fuente = '{FUENTE}') <> {len(renglones)} THEN
        RAISE EXCEPTION 'Algún producto del Excel no se encontró en el menú';
    END IF;
    IF EXISTS (SELECT 1 FROM historial_ventas h JOIN variantes v ON v.producto_id = h.producto_id
                WHERE h.fuente = '{FUENTE}' AND h.variante_id IS NULL) THEN
        RAISE EXCEPTION 'Un producto por tamaño quedó sin tamaño';
    END IF;
END $$;

-- Lo que llevan los combos (para calcular insumos): la primera opción de cada grupo.
INSERT INTO historial_ventas (dia, producto_id, variante_id, cantidad, componente, fuente)
SELECT h.dia, o.producto_id, o.variante_id, sum(h.cantidad * g.cantidad), true, h.fuente
  FROM historial_ventas h
  JOIN combo_grupos g ON g.combo_id = h.producto_id
 CROSS JOIN LATERAL (SELECT producto_id, variante_id FROM combo_opciones
                      WHERE grupo_id = g.id ORDER BY id LIMIT 1) o
 WHERE h.fuente = '{FUENTE}' AND NOT h.componente
 GROUP BY h.dia, o.producto_id, o.variante_id, h.fuente;
COMMIT;""")


if __name__ == "__main__":
    main()
